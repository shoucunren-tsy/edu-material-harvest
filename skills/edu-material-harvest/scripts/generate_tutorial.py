# -*- coding: utf-8 -*-
"""生成使用教程 docx：使用教程_高校公开资料采集.docx

面向普通用户：快速上手 / 标准工作流 / 配置说明 / 常见问题。不写技术实现细节。
页脚必须有页码（居中）。
"""
import os
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.stderr.write(
        "缺少 python-docx，无法生成使用教程 docx。\n"
        "  解决：pip install python-docx\n")
    sys.exit(3)

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
OUT = os.path.join(SKILL_DIR, "使用教程_高校公开资料采集.docx")

CJK = "微软雅黑"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = CJK
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10.5, bold=False, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_font(r, size=size)
    return p


def heading(doc, text, size=15):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=size, bold=True, color=(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_page_number_footer(doc):
    """页脚居中页码：第 X 页 / 共 Y 页。"""
    sec = doc.sections[0]
    footer = sec.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def field(instr):
        r = p.add_run()
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        r._r.append(f1); r._r.append(it); r._r.append(f2)
        set_font(r, size=9)

    r0 = p.add_run("第 "); set_font(r0, size=9)
    field("PAGE")
    r1 = p.add_run(" 页 / 共 "); set_font(r1, size=9)
    field("NUMPAGES")
    r2 = p.add_run(" 页"); set_font(r2, size=9)


def build():
    doc = Document()
    # 正文默认字体
    style = doc.styles["Normal"]
    style.font.name = CJK
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    para(doc, "高校公开资料采集 · 使用教程", size=20, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    para(doc, "从官方免登录公开源自动发现并采集培养方案 / 专业介绍 / 课程大纲 / 年度报告，"
              "逐份落盘并登记可溯源的来源链接", size=10.5,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

    # 一、这是什么
    heading(doc, "一、这是什么")
    para(doc, "给定「学校名 + 官方域名」，自动从**官方公开网页**发现并采集四类材料：培养方案、"
              "专业介绍、课程大纲、年度报告（就业质量 / 本科教学质量）。每采一份，都会记下它的"
              "**下载直链**和**挂链接的那个网页**，最后汇总成一份可点击、可核对的台账。")
    para(doc, "**适用于任意一所中国高校（*.edu.cn），不限地域。**"
              "核心价值：结果**可溯源、可复现**——你要做分析时，能一键跳回学校官网核对原文。")

    # 二、能采什么、不能采什么
    heading(doc, "二、能采什么、不能采什么")
    para(doc, "能采（免登录公开源）：")
    bullet(doc, "学校官网、教务处、招生网、信息公开网上公开挂出的 PDF / 打包件（RAR、ZIP）。")
    bullet(doc, "只有网页正文、没有附件的材料（如部分专业介绍、通识课大纲）——只登记链接，不下载。")
    para(doc, "采不到（别抱期望）：")
    bullet(doc, "登录后才能看的教务系统内材料（本工具不登录、不用账号）。")
    bullet(doc, "没有任何一所学校把全校课程大纲公开挂在网上；整校批量大纲多在教务系统里，采不到。")
    bullet(doc, "被网站反爬（WAF / JS 挑战）挡住的页面——这是真实边界，工具会如实记下，不硬闯。")

    # 三、快速上手
    heading(doc, "三、快速上手（三步）")
    para(doc, "第 1 步 · 给两个信息：**采哪所学校**、**它的官方域名**。"
              "域名由 Claude 先用联网搜索确认为官方 *.edu.cn（**绝不瞎猜**）。")
    para(doc, "第 2 步 · 自动发现（只侦察、不下载）：Claude 会摸清该校材料挂在哪些栏目页、"
              "附件直链是什么，并对每条候选**实测**可达性与文件类型，产出一份“候选确认表”给你过目。"
              "如果目标材料挂在被反爬挡住的子站、自动爬不到，Claude 会改用“先联网搜索定位到挂着材料的那一页，"
              "再交给脚本枚举下载”的精准方式——两种方式产出的台账完全一样。")
    para(doc, "第 3 步 · 一键采集 + 出台账：确认无误后，自动下载、校验完整性、按内容去重、"
              "分目录归档，并生成 Excel 台账与来源链接台账。")
    para(doc, "全程只需你给“哪所学校、哪类材料”，其余交给 Claude。", bold=True)

    # 四、标准工作流
    heading(doc, "四、标准工作流")
    for i, s in enumerate([
        "明确院校与材料类型（培养方案 / 专业介绍 / 课程大纲 / 年度报告）。",
        "解析域名：联网确认该校官方 *.edu.cn 域名（不猜）。",
        "自动发现：探明栏目页与附件直链，每条实测可达性与文件类型（只读，不下载）。",
        "看候选确认表：每条候选都带 HTTP 状态 / 内容类型 / 字节数，是“真链接”的证据。",
        "一键采集：下载 → 校验文件头 → 与已有资料按内容去重 → 分目录落盘。",
        "打包件（RAR/ZIP）：自动解压，按内层文件夹重组，逐份登记。",
        "命名：培养方案/专业介绍按“年份_专业”；课程大纲按“课程号_课程名”。",
        "出台账：Excel（5 张表）+ 2 个 CSV，逐条含来源直链与承载页。",
        "收尾：临时文件走回收站；把新发现的学校来源记入来源清单，下次不用重找。",
    ], 1):
        p = doc.add_paragraph()
        r = p.add_run("%d. %s" % (i, s)); set_font(r)
        p.paragraph_format.space_after = Pt(3)

    # 五、产物长什么样
    heading(doc, "五、产物长什么样")
    para(doc, "每个文件都在目录里有一行记录，字段固定：")
    bullet(doc, "学校 | 材料类型 | 专业/说明 | 年份 | 文件名 | 大小 | 来源直链 | 承载页 | 采集日期")
    para(doc, "另有一张“链接台账”，专门登记页面级链接（哪怕没有附件也会登记），"
              "让你随时能跳回官方页面。")
    para(doc, "Excel 台账共 5 张表：总览 / 按院校汇总 / 明细 / 链接台账 / 未成功入库；"
              "另出 2 个 CSV（采集清单、链接台账）。")
    para(doc, "目录结构：新增数据 / <学校> / <材料类型> / <文件>")

    # 六、配置说明
    heading(doc, "六、配置说明")
    bullet(doc, "交付目录：由你指定，脚本按“学校 / 材料类型”自动建子目录。")
    bullet(doc, "RAR 解压：自动探测本机已有的 7-Zip / Bandizip，无需你安装任何软件。")
    bullet(doc, "外部工具（解压器 / curl / PowerShell）全部自动探测，也可用环境变量覆盖。")
    bullet(doc, "原始数据：你的原始资料目录只读，本工具绝不改动、删除、改名。")
    bullet(doc, "删除：所有临时文件都进回收站，不做永久删除；无回收站时改为隔离，文件不丢。")

    # 七、常见问题
    heading(doc, "七、常见问题")
    qa = [
        ("为什么某校某类材料一份都没采到？",
         "多半是该校把材料放在登录后才可见的教务系统里，或近年已停止公开。工具只采免登录公开源，"
         "采不到会如实记入“未成功入库”，不会编造。"),
        ("有的材料只给了链接、没下文件？",
         "有些材料（如专业介绍、通识课大纲）只有网页正文、没有附件。按约定只登记链接，不下载网页文件，"
         "台账里照样能点击跳转。"),
        ("重复的文件会怎样？",
         "按文件内容自动去重，和已有资料重复的会被剔除并记录为“重复”，不会重复入库。"
         "同一批再跑一遍也不会重复下载（已存在的会自动跳过）。"),
        ("能采任意学校吗？",
         "适用于任何 *.edu.cn 高校，不限地域。但“有多少公开”因校而异——先跑一轮自动发现，"
         "看候选确认表就知道实际能采到多少。"),
        ("为什么会采到“看起来不像材料”的网页？",
         "极少数情况下，新闻页或招生页正文里出现了“培养方案”“年度报告”等词，可能被误判。工具已内置"
         "栏目黑名单与类型复核来尽量剔除；若仍发现不对的条目，告诉 Claude 用“指定承载页”方式重采即可，"
         "不会影响已确认的文件。"),
        ("有些链接提示被拦 / 打不开？",
         "那是网站有反爬（WAF / 登录墙 / JS 挑战），属于真实边界，不是工具坏了。"
         "工具会如实记下“被拦”，不会伪造数据，也不会硬闯。"),
        ("链接会不会失效？",
         "高校栏目常改版。台账里每条链接都是采集当天实测可达的；用之前建议再核一次。"),
        ("能在其它 Agent（Codex、WorkBuddy、千问办公）里用吗？",
         "能。本 skill 带一个零依赖的 MCP 服务，装一次即可让这些客户端用同一套采集能力"
         "（见第八节）。不需要你额外安装 Python 包。"),
    ]
    for q, a in qa:
        para(doc, "Q：" + q, bold=True, space_after=2)
        para(doc, "A：" + a, space_after=8)

    # 八、在别的 Agent 里用
    heading(doc, "八、在别的 Agent 里用（Claude Code / Codex / WorkBuddy / 千问办公）")
    para(doc, "这套采集能力不止 Claude Code 能用。它带一个**零依赖**的 MCP 服务，"
              "把同一套本事交给别的 Agent，**你不需要额外安装任何东西**。")
    para(doc, "接入三步（以 Claude Code 为例）：")
    for i, s in enumerate([
        "在 skill 目录下运行：python scripts/mcp_server.py --print-config，"
        "它会显示两行路径（解释器和脚本）。",
        "接入：python scripts/mcp_server.py --install claude（Codex 用 --install codex，"
        "WorkBuddy / 千问办公 用第一条打印出的信息在它们的「连接器 / MCP」里新建）。",
        "在客户端里就能看到 harvest_school（一键采集）、harvest_pages（指定承载页采集）"
        "等工具；说「采集 XX 大学的培养方案」即可。",
    ], 1):
        p = doc.add_paragraph()
        r = p.add_run("%d. %s" % (i, s)); set_font(r)
        p.paragraph_format.space_after = Pt(3)
    bullet(doc, "长任务会自动转后台：工具先返回一个任务号，稍等片刻再取结果，"
                "不会把客户端卡到超时。")
    bullet(doc, "想让它更省心：把「某校挂着材料的那一页」网址直接告诉它，它就走精准直采，"
                "比全站找更稳。")
    para(doc, "注意：命令里的 Python 路径要与实际一致；若客户端找不到 Python，"
              "就用 --print-config 打印的绝对路径。", size=10.5)

    add_page_number_footer(doc)
    doc.save(OUT)
    print("已生成：", OUT)


if __name__ == "__main__":
    build()
