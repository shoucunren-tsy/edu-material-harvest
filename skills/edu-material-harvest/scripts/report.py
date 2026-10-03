# -*- coding: utf-8 -*-
"""高校公开资料采集 · 台账生成

用法：
  python report.py records.json --out <记录目录> [--html-links html_links.json]
      [--date YYYY-MM-DD] [--csv-only]

输入：
  records.json     download.py 产出的记录（含 status=新增/重复/失败）
  html_links.json  可选。纯 HTML 来源页（无附件）的登记：
                   [{school, type, form, link, desc}]   —— 只登记链接、不下载

输出（<out> 目录）：
  采集清单.csv     逐文件明细
  链接台账.csv     页面级链接台账（承载页 + HTML 来源页）
  采集台账.xlsx    多表：总览 / 按院校汇总 / 明细 / 链接台账 / 未成功入库
                   （缺 openpyxl 时默认报错退出；加 --csv-only 降级为只出 CSV）
"""
import os
import sys
import csv
import argparse
from datetime import date as _date

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C  # noqa: E402

DATE_DEFAULT = None   # None = 用今天（不写死日期，可移植）

DETAIL_COLS = ["序号", "学校", "材料类型", "专业/说明", "年份", "文件名", "大小(MB)",
               "来源直链", "承载页", "采集日期"]
LEDGER_COLS = ["序号", "学校", "来源类型", "形式", "链接", "说明/承载内容", "采集日期"]
FAIL_COLS = ["序号", "学校", "材料类型", "专业/说明", "年份", "结果", "来源直链", "承载页"]


def _ok(records):
    """入库成功的记录（有文件名）。"""
    return [r for r in records if r.get("file")]


def build_ledger(records, html_links, date):
    """页面级链接台账：已下文件的承载页（去重）+ HTML 来源页。"""
    rows, seen = [], set()
    for r in _ok(records):
        pg = (r.get("page") or "").strip()
        if not pg or pg in seen:
            continue
        seen.add(pg)
        rows.append([r.get("school", ""), "教务处/学院栏目", "承载页", pg,
                     (r.get("desc") or "")[:40], date])
    for h in (html_links or []):
        rows.append([h.get("school", ""), h.get("type", ""), h.get("form", "HTML正文页"),
                     h.get("link", ""), h.get("desc", ""), date])
    for i, row in enumerate(rows, 1):
        row.insert(0, i)
    return rows


def write_csvs(records, ledger, outdir, date, fail):
    p1 = os.path.join(outdir, "采集清单.csv")
    with open(p1, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(DETAIL_COLS)
        for i, r in enumerate(_ok(records), 1):
            w.writerow([i, r.get("school", ""), r.get("kind", ""), r.get("desc", ""),
                        r.get("year", ""), r.get("file", ""), r.get("size_mb", ""),
                        r.get("url", ""), r.get("page", ""), date])
    p2 = os.path.join(outdir, "链接台账.csv")
    with open(p2, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(LEDGER_COLS)
        w.writerows(ledger)
    return p1, p2


def write_xlsx(records, ledger, outdir, date, csv_only=False):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
    except ImportError:
        # xlsx 是核心产物：缺库要吵不要哑。--csv-only 时才算有意跳过。
        msg = ("未安装 openpyxl，无法生成 Excel 台账（xlsx 是本工具的核心产物）。\n"
               "  解决：pip install openpyxl\n"
               "  或加 --csv-only 只出 CSV（明细与链接台账均不丢，只是没有 Excel 多表汇总）。")
        if csv_only:
            print("WARN:", msg)
            return None
        print("ERROR:", msg, file=sys.stderr)
        sys.exit(3)
    from collections import Counter, defaultdict

    wb = openpyxl.Workbook()
    hf = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDEBF7")

    def hd(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = hf, fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    del wb["Sheet"]

    # 总览
    ws = wb.create_sheet("总览")
    hd(ws, ["材料类型", "份数"])
    cnt = Counter(r.get("kind", "") for r in _ok(records))
    for k, v in sorted(cnt.items()):
        ws.append([k, v])
    ws.append(["合计", sum(cnt.values())])

    # 按院校汇总
    ws = wb.create_sheet("按院校汇总")
    hd(ws, ["学校", "材料类型", "份数"])
    agg = defaultdict(lambda: Counter())
    for r in _ok(records):
        agg[r.get("school", "")][r.get("kind", "")] += 1
    for school in sorted(agg):
        for kind, n in sorted(agg[school].items()):
            ws.append([school, kind, n])

    # 明细
    ws = wb.create_sheet("明细")
    hd(ws, DETAIL_COLS)
    for i, r in enumerate(_ok(records), 1):
        ws.append([i, r.get("school", ""), r.get("kind", ""), r.get("desc", ""),
                   r.get("year", ""), r.get("file", ""), r.get("size_mb", ""),
                   r.get("url", ""), r.get("page", ""), date])

    # 链接台账
    ws = wb.create_sheet("链接台账")
    hd(ws, LEDGER_COLS)
    for row in ledger:
        ws.append(row)

    # 未成功入库
    ws = wb.create_sheet("未成功入库")
    hd(ws, FAIL_COLS)
    for i, r in enumerate([r for r in records if not r.get("file")], 1):
        ws.append([i, r.get("school", ""), r.get("kind", ""), r.get("desc", ""),
                   r.get("year", ""), r.get("status", ""), r.get("url", ""), r.get("page", "")])

    xp = os.path.join(outdir, "采集台账.xlsx")
    wb.save(xp)
    return xp


def main():
    ap = argparse.ArgumentParser(description="高校公开资料采集·台账生成")
    ap.add_argument("records", help="records.json")
    ap.add_argument("--out", required=True, help="台账输出目录")
    ap.add_argument("--html-links", help="纯 HTML 来源页登记 json（可选）")
    ap.add_argument("--date", default=DATE_DEFAULT, help="采集日期（缺省=今天）")
    ap.add_argument("--csv-only", action="store_true",
                    help="只出 CSV，不生成 xlsx（无 openpyxl 时的降级路径）")
    a = ap.parse_args()

    the_date = a.date or _date.today().isoformat()
    records = C.load_json(a.records, [])
    html_links = C.load_json(a.html_links, []) if a.html_links else []
    os.makedirs(a.out, exist_ok=True)

    ledger = build_ledger(records, html_links, the_date)
    p1, p2 = write_csvs(records, ledger, a.out, the_date, None)
    xp = None if a.csv_only else write_xlsx(records, ledger, a.out, the_date)

    ok = len(_ok(records))
    print("明细 %d 份；链接台账 %d 行" % (ok, len(ledger)))
    wrote = [p1, p2] + ([xp] if xp else [])
    print("已写：", "\n      ".join(wrote))


if __name__ == "__main__":
    main()
